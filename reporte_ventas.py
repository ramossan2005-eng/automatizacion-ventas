"""
Automatización de reportes de ventas.

Lee todos los archivos de Excel de la carpeta datos/, los limpia,
los une en una sola tabla y genera reporte_ventas.xlsx con:
  - Resumen general
  - Ventas por mes, por producto, por ciudad y por vendedor
  - Gráficas nativas de Excel
  - Los datos limpios, por si el cliente quiere revisarlos

Uso:  python3 reporte_ventas.py
"""

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

CARPETA_DATOS = Path("datos")
ARCHIVO_SALIDA = Path("reporte_ventas.xlsx")

# Cada archivo puede llamar distinto a la misma columna.
# Aquí decimos a qué nombre estándar corresponde cada variante.
NOMBRES_COLUMNAS = {
    "fecha": "fecha",
    "fecha de venta": "fecha",
    "producto": "producto",
    "ciudad": "ciudad",
    "vendedor": "vendedor",
    "cantidad": "cantidad",
    "cant.": "cantidad",
    "precio": "precio",
    "precio unitario": "precio",
    "valor": "precio",
}

MESES = {1: "Enero", 2: "Febrero", 3: "Marzo", 4: "Abril", 5: "Mayo", 6: "Junio",
         7: "Julio", 8: "Agosto", 9: "Septiembre", 10: "Octubre",
         11: "Noviembre", 12: "Diciembre"}


# ---------------------------------------------------------------- 1. Leer
def leer_archivos(carpeta):
    archivos = sorted(carpeta.glob("*.xlsx"))
    if not archivos:
        raise FileNotFoundError(f"No hay archivos .xlsx en la carpeta {carpeta}/")

    tablas = []
    for archivo in archivos:
        df = pd.read_excel(archivo)
        # Estandarizar nombres de columna: minúsculas, sin espacios extra
        df.columns = [str(c).strip().lower() for c in df.columns]
        df = df.rename(columns=NOMBRES_COLUMNAS)
        df["archivo_origen"] = archivo.name
        tablas.append(df)
        print(f"  Leído: {archivo.name} ({len(df)} filas)")
    return pd.concat(tablas, ignore_index=True)


# -------------------------------------------------------------- 2. Limpiar
def convertir_precio(valor):
    """Convierte '$180.000' o 180000 en el número 180000."""
    if pd.isna(valor):
        return None
    if isinstance(valor, str):
        valor = valor.replace("$", "").replace(".", "").replace(",", "").strip()
    return float(valor)


def limpiar(df):
    filas_iniciales = len(df)
    reporte = {}

    # Textos: quitar espacios y unificar mayúsculas ("  BOGOTÁ " -> "Bogotá")
    for col in ["producto", "ciudad", "vendedor"]:
        df[col] = df[col].astype("string").str.strip().str.capitalize()

    df["precio"] = df["precio"].apply(convertir_precio)
    df["fecha"] = pd.to_datetime(df["fecha"], errors="coerce")
    df["cantidad"] = pd.to_numeric(df["cantidad"], errors="coerce")

    # Duplicados exactos (sin contar el archivo de origen)
    columnas_clave = ["fecha", "producto", "ciudad", "vendedor", "cantidad", "precio"]
    duplicados = df.duplicated(subset=columnas_clave).sum()
    df = df.drop_duplicates(subset=columnas_clave)
    reporte["Filas duplicadas eliminadas"] = int(duplicados)

    # Filas sin datos esenciales
    incompletas = df[columnas_clave].isna().any(axis=1)
    reporte["Filas incompletas eliminadas"] = int(incompletas.sum())
    df = df[~incompletas].copy()

    df["cantidad"] = df["cantidad"].astype(int)
    df["total"] = df["cantidad"] * df["precio"]
    df["mes_num"] = df["fecha"].dt.month
    df["mes"] = df["mes_num"].map(MESES)

    reporte["Filas leídas"] = filas_iniciales
    reporte["Filas válidas"] = len(df)
    return df.sort_values("fecha"), reporte


# ------------------------------------------------------------- 3. Resumir
def resumir(df):
    por_mes = (df.groupby(["mes_num", "mes"], as_index=False)
                 .agg(unidades=("cantidad", "sum"), ventas=("total", "sum"))
                 .sort_values("mes_num").drop(columns="mes_num"))

    def agrupar(col):
        return (df.groupby(col, as_index=False)
                  .agg(unidades=("cantidad", "sum"), ventas=("total", "sum"))
                  .sort_values("ventas", ascending=False))

    return {
        "Por mes": por_mes,
        "Por producto": agrupar("producto"),
        "Por ciudad": agrupar("ciudad"),
        "Por vendedor": agrupar("vendedor"),
    }


# ------------------------------------------------------------ 4. Exportar
def exportar(df, resumenes, reporte_limpieza, salida):
    resumen_general = pd.DataFrame({
        "Indicador": [
            "Ventas totales",
            "Unidades vendidas",
            "Número de ventas",
            "Venta promedio",
            "Producto con más ventas",
            "Mejor ciudad",
            *reporte_limpieza.keys(),
        ],
        "Valor": [
            df["total"].sum(),
            df["cantidad"].sum(),
            len(df),
            df["total"].mean(),
            resumenes["Por producto"].iloc[0]["producto"],
            resumenes["Por ciudad"].iloc[0]["ciudad"],
            *reporte_limpieza.values(),
        ],
    })

    datos_limpios = df[["fecha", "mes", "producto", "ciudad", "vendedor",
                        "cantidad", "precio", "total", "archivo_origen"]]

    with pd.ExcelWriter(salida, engine="openpyxl") as writer:
        resumen_general.to_excel(writer, sheet_name="Resumen", index=False)
        for nombre, tabla in resumenes.items():
            tabla.to_excel(writer, sheet_name=nombre, index=False)
        datos_limpios.to_excel(writer, sheet_name="Datos limpios", index=False)

    dar_formato(salida)


def dar_formato(salida):
    wb = load_workbook(salida)
    encabezado = PatternFill("solid", fgColor="0F4C5C")
    pesos = '"$"#,##0'

    for ws in wb.worksheets:
        # Encabezados
        for celda in ws[1]:
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = encabezado
            celda.alignment = Alignment(horizontal="center")
            celda.value = str(celda.value).replace("_", " ").capitalize()
        # Ancho de columnas
        for i, col in enumerate(ws.columns, start=1):
            largo = max(len(str(c.value or "")) for c in col)
            ws.column_dimensions[get_column_letter(i)].width = min(largo + 4, 40)
        # Formato de pesos
        for col in ws.iter_cols(min_row=1):
            if col[0].value in ("Ventas", "Precio", "Total"):
                for c in col[1:]:
                    c.number_format = pesos
            if col[0].value == "Fecha":
                for c in col[1:]:
                    c.number_format = "DD/MM/YYYY"
        ws.freeze_panes = "A2"

    # Formato de la hoja Resumen
    ws = wb["Resumen"]
    for fila in (2, 5):  # Ventas totales y venta promedio
        ws.cell(fila, 2).number_format = pesos
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 22

    # Gráficas
    agregar_grafica(wb["Por mes"], "Ventas por mes")
    for hoja in ("Por producto", "Por ciudad", "Por vendedor"):
        agregar_grafica(wb[hoja], f"Ventas {hoja.lower()}")

    wb.save(salida)


def agregar_grafica(ws, titulo):
    grafica = BarChart()
    grafica.title = titulo
    grafica.y_axis.title = "Ventas (COP)"
    grafica.y_axis.numFmt = '"$"#,##0'
    grafica.legend = None
    grafica.x_axis.delete = False
    grafica.y_axis.delete = False

    grafica.height, grafica.width = 8, 16

    datos = Reference(ws, min_col=3, min_row=1, max_row=ws.max_row)
    categorias = Reference(ws, min_col=1, min_row=2, max_row=ws.max_row)
    grafica.add_data(datos, titles_from_data=True)
    grafica.set_categories(categorias)
    ws.add_chart(grafica, "E2")


# ------------------------------------------------------------------ Main
def main():
    print("1. Leyendo archivos...")
    df = leer_archivos(CARPETA_DATOS)

    print("2. Limpiando datos...")
    df, reporte_limpieza = limpiar(df)
    for k, v in reporte_limpieza.items():
        print(f"  {k}: {v}")

    print("3. Calculando resúmenes...")
    resumenes = resumir(df)

    print("4. Generando reporte...")
    exportar(df, resumenes, reporte_limpieza, ARCHIVO_SALIDA)
    print(f"\nListo: {ARCHIVO_SALIDA} ({df['total'].sum():,.0f} COP en ventas)")


if __name__ == "__main__":
    main()
